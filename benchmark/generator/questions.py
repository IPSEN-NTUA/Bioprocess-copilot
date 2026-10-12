"""
Benchmark question set for the pilot corpus.

Every gold answer is computed by reading the CORPUS files only (never the generator or the
answer key), so each answer is demonstrably derivable from what a system under test can see.

    python questions.py [--pilot ../pilot]   ->  pilot/questions/questions.{json,csv,md}

Tiers: L1 single-document lookup, L2 single document + small calculation, L3 join across 2-3
documents, L4 provenance / fault diagnosis, L5 campaign-level synthesis, U unanswerable or
false premise (the correct behaviour is to say so).
"""
from __future__ import annotations

import argparse
import datetime as dtm
import glob
import json
import re
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd


class Corpus:
    def __init__(self, root: Path):
        self.root = root
        self._tr = None

    def p(self, rel):
        return self.root / rel

    def text(self, rel):
        return self.p(rel).read_text()

    # ---- HPLC (TR) ----------------------------------------------------------
    def tr(self) -> pd.DataFrame:
        if self._tr is None:
            rows = []
            for f in sorted(glob.glob(str(self.p("hplc/processed/*.xlsx")))):
                it = openpyxl.load_workbook(f, data_only=True)["Results"].iter_rows(values_only=True)
                h = next(it)
                rows += [dict(zip(h, r)) for r in it]
            self._tr = pd.DataFrame(rows)
        return self._tr

    def tr_value(self, sample, analyte, rerun=True):
        d = self.tr()
        d = d[(d.SampleID == sample) & (d.Analyte == analyte)]
        if rerun:
            r = d[d.Comment.fillna("").str.startswith("re-run")]
            if len(r):
                return r.Result_gL.iloc[0]
        return d[~d.Comment.fillna("").str.startswith("re-run")].Result_gL.iloc[0]

    def broth_samples(self, run):
        reg = self.lims()
        return list(reg[(reg.run_id == run) & (reg.type == "broth")].sample_id)

    def final(self, run, analyte="Xylitol"):
        return self.tr_value(self.broth_samples(run)[-1], analyte)

    # ---- other sources --------------------------------------------------------
    def lims(self):
        return pd.read_csv(self.p("lims/sample_register.csv"), comment="#")

    def xlsx_sheet(self, rel, sheet, header_row=1):
        ws = openpyxl.load_workbook(self.p(rel), data_only=True)[sheet]
        rows = list(ws.iter_rows(values_only=True))
        return rows

    def ev(self, run):
        rows = self.xlsx_sheet(f"spreadsheets/EV_{run}.xlsx", "Offline data")
        df = pd.DataFrame(rows[1:], columns=rows[0])
        calc = dict(self.xlsx_sheet(f"spreadsheets/EV_{run}.xlsx", "Calculated"))
        info = dict(r[:2] for r in self.xlsx_sheet(f"spreadsheets/EV_{run}.xlsx", "Run info"))
        return df, calc, info

    def md(self, sheet):
        rows = self.xlsx_sheet("spreadsheets/MD_fermentation_master_2025.xlsx", sheet)
        return pd.DataFrame(rows[1:], columns=rows[0])

    def registry(self):
        rows = self.xlsx_sheet("strains/strain_registry.xlsx", "Strains")
        return pd.DataFrame(rows[1:], columns=rows[0]).set_index("Strain ID")

    def online(self, run):
        f = glob.glob(str(self.p(f"online/{run}_*.txt"))) or glob.glob(str(self.p(f"online/{run.replace('-', '')}_*.csv")))
        f = f[0]
        if f.endswith(".txt"):
            lines = [l for l in open(f) if not l.startswith("#")]
            inoc = next(l for l in lines if l.startswith("Inoculation")).split(":", 1)[1].strip()
            inoc = dtm.datetime.strptime(inoc, "%d %b %Y %H:%M:%S")
            i = next(k for k, l in enumerate(lines) if l.startswith("Date"))
            hdr = lines[i].strip().split("\t")
            df = pd.DataFrame([l.rstrip("\r\n").split("\t") for l in lines[i + 1:]], columns=hdr)
            ts = pd.to_datetime(df.Date + " " + df.Time, format="%d.%m.%Y %H:%M:%S")
            out = pd.DataFrame(dict(t=(ts - inoc).dt.total_seconds() / 3600, pO2=df["pO2[%]"].astype(float),
                                    stirrer=df["Stirrer[rpm]"].astype(float)))
        else:
            df = pd.read_csv(f, comment="#")
            k = df.index[df.AGIT_rpm > 400][0]                       # agitation step = inoculation
            out = pd.DataFrame(dict(t=df.Elapsed_h - df.Elapsed_h[k], pO2=df.DO_pct, stirrer=df.AGIT_rpm))
        return out[out.t >= 0].reset_index(drop=True), Path(f).relative_to(self.root.parent)

    def coa(self, lot):
        t = self.text(f"materials/CoA_{lot}.md")
        return {m.group(1).strip(): float(m.group(2)) for m in re.finditer(r"\| ([^|]+) \| ([\d.]+) \|", t)}

    def notes(self):
        return {Path(f).stem: Path(f).read_text() for f in glob.glob(str(self.p("notes/FB-*.md")))}


def ev_path(r):
    return f"corpus/spreadsheets/EV_{r}.xlsx"


MD = "corpus/spreadsheets/MD_fermentation_master_2025.xlsx"
REG = "corpus/strains/strain_registry.xlsx"


def run_table(C: Corpus) -> pd.DataFrame:
    """Operator, vessel and strain of every bioreactor run, from the documents that record them:
    EV workbooks (Run info), MD master log (Runs sheet) and PN notes."""
    rows = []
    for f in sorted(glob.glob(str(C.p("spreadsheets/EV_FB-*.xlsx")))):
        info = dict(r[:2] for r in C.xlsx_sheet(Path(f).relative_to(C.root), "Run info"))
        m = re.match(r"(YX-\d+) \(clone (c\d)\)", info["Strain"])
        rows.append(dict(run=info["Run"], operator="EV", vessel=info["Vessel"].split(" ")[0], strain=m.group(1),
                         clone=m.group(2), purpose=info["Purpose"]))
    for _, r in C.md("Runs").iterrows():
        rows.append(dict(run=r.Run.replace("FB", "FB-0"), operator="MD", vessel=r.Vsl, strain=r.Strain, clone=None,
                         purpose=None))
    for k, v in C.notes().items():
        if k.endswith("_PN"):
            m = re.search(r"Ran (YX-\d+) (c\d) in (R\d)\.", v)
            rows.append(dict(run=k.split("_")[0], operator="PN", vessel=m.group(3), strain=m.group(1), clone=m.group(2),
                             purpose=None))
    return pd.DataFrame(rows).sort_values("run").reset_index(drop=True)


def build_questions(C: Corpus) -> list[dict]:
    Q = []

    def q(tier, category, question, answer, evidence, grading, value=None, tol=None, key_points=None, hops=1):
        Q.append(dict(id=f"{tier}-{sum(1 for x in Q if x['tier'] == tier) + 1:02d}", tier=tier, category=category,
                      question=question, gold_answer=answer, value=value, tolerance=tol, grading=grading,
                      key_points=key_points or [], evidence=evidence, hops=hops))

    tr = C.tr()
    reg = C.registry()
    mdr, mds, mdm = C.md("Runs"), C.md("Samples"), C.md("Media & feeds")
    notes = C.notes()
    RT = run_table(C)
    assert RT.run.is_unique and len(RT) == 25, "every run must be documented exactly once"
    minutes = {Path(f).stem.split("_")[1]: Path(f).read_text() for f in glob.glob(str(C.p("meetings/*.md")))}
    names = {n.strip() for t_ in minutes.values() for n in re.search(r"\*\*Present:\*\* (.+)", t_).group(1).split(",")}
    NAME = {"".join(w[0] for w in n.split()): n for n in names}                    # initials -> full name

    def md_final(run, col="Xol"):
        return mds[mds.Run == run.replace("FB-0", "FB")][col].iloc[-1]

    # ======================= L1: single-document lookup ======================
    s = mdr[mdr.Run == "FB13"].Strain.iloc[0]
    q("L1", "lookup", "Which strain was used in bioreactor run FB-013?", s, [MD, "corpus/notes/FB-013_MD.md"], "exact", value=s)
    v = re.search(r"in (R\d)\.", notes["FB-022_PN"]).group(1)
    q("L1", "lookup", "In which bioreactor vessel was FB-022 run?", v, ["corpus/notes/FB-022_PN.md"], "exact", value=v)
    f = C.coa("C5-2503")["Formic acid"]
    q("L1", "lookup", "What formic acid concentration does the certificate of analysis report for C5 lot C5-2503?",
      f"{f:.2f} g/L", ["corpus/materials/CoA_C5-2503.md"], "numeric", value=f, tol=0.01)
    by = reg.loc["YX-004", "By"]
    q("L1", "lookup", "Who constructed strain YX-004?", by, [REG, "corpus/strains/construction_YX-004.md"], "exact", value=by)
    eff = re.search(r"\*\*Effective from:\*\* (\S+)", C.text("sops/SOP-ANA-003_v2.md")).group(1)
    q("L1", "lookup", "From which date is version 2 of the HPLC method SOP-ANA-003 effective?", eff,
      ["corpus/sops/SOP-ANA-003_v2.md"], "exact", value=eff)
    m = re.search(r"Start the feed at ([\d.]+) h.*?at ([\d.]+) mL/h", C.text("sops/SOP-FB-001_v3.md"), re.S)
    q("L1", "lookup", "According to the fed-batch SOP, when is the feed started and at what rate?",
      f"At {m.group(1)} h, at {m.group(2)} mL/h", ["corpus/sops/SOP-FB-001_v3.md"], "key_points",
      key_points=[f"{m.group(1)} h", f"{m.group(2)} mL/h"])
    c6 = reg.loc["YX-006", "Selected clone"]
    q("L1", "lookup", "Which clone of YX-006 is recorded as selected in the strain registry?", c6, [REG], "exact", value=c6)
    _, calc, _ = C.ev("FB-001")
    x = calc["Final xylitol (g/L)"]
    q("L1", "lookup", "What final xylitol concentration did Elena Varga's workbook report for FB-001?", f"{x:.1f} g/L",
      [ev_path("FB-001")], "numeric", value=x, tol=0.1)
    dl = int(tr[tr.SampleID == "FB018-S06"].Dilution.iloc[0])
    q("L1", "lookup", "What dilution factor was used for the HPLC analysis of sample FB018-S06?", f"1:{dl}",
      ["corpus/hplc/processed/HPLC_results_2025-06.xlsx", "corpus/hplc/raw/"], "exact", value=dl)

    # ======================= L2: one document + small calculation ============
    t3 = tr[(tr.Run == "FB-003") & (tr.Analyte == "Glucose") & tr.SampleID.str.contains("-S")].copy()
    t3["h"] = t3.SampleID.map(dict(zip(C.lims().sample_id, C.lims().process_time_h)))
    t_glc = t3[(t3.h > 0.2) & (t3.Flag == "<LOD")].h.min()
    df3, _, _ = C.ev("FB-003")
    assert pd.isna(df3[df3["Time (h)"] == t_glc]["Glucose (g/L)"].iloc[0])        # EV workbook agrees (blank)
    q("L2", "aggregation", "In FB-003, at which sampling time was glucose first below the HPLC detection limit?",
      f"{t_glc:g} h", [ev_path("FB-003"), "corpus/hplc/processed/HPLC_results_2025-02.xlsx"], "numeric",
      value=float(t_glc), tol=0)
    df, _, _ = C.ev("FB-001")
    br = df[df.Sample.str.contains("-S")].copy()
    br["E"] = br["Ethanol (g/L)"].fillna(0)
    i = br.E.idxmax()
    q("L2", "aggregation", "What was the peak ethanol concentration in FB-001, and at which sample time?",
      f"{br.E[i]:.2f} g/L at {br['Time (h)'][i]:g} h", [ev_path("FB-001")], "key_points",
      value=float(br.E[i]), tol=0.02, key_points=[f"{br.E[i]:.2f} g/L", f"{br['Time (h)'][i]:g} h"])
    b = float(re.search(r"Base used: ([\d.]+) mL", notes["FB-008_EV"]).group(1))
    q("L2", "lookup", "How much base was used in FB-008 according to the lab note?", f"{b:.1f} mL",
      ["corpus/notes/FB-008_EV.md"], "numeric", value=b, tol=0.1)
    _, calc, _ = C.ev("FB-018")
    q("L2", "aggregation", "What was the maximum dry weight reached in FB-018?", f"{calc['Max DW (g/L)']:.2f} g/L",
      [ev_path("FB-018")], "numeric", value=calc["Max DW (g/L)"], tol=0.01)
    t = C.text("notes/SF_screen_P2_PN.md")
    hr = re.search(r"Xylitol \(g/L, HPLC\) at (\d+) h", t).group(1)
    mm = re.search(r"\| YX-002 c3 \| ([\d.]+) \| (\d+) \|", t)
    q("L2", "aggregation", f"In the March shake-flask clone screen, what mean xylitol did YX-002 clone c3 reach at {hr} h, "
      "and from how many flasks?",
      f"{mm.group(1)} g/L from {mm.group(2)} flasks", ["corpus/notes/SF_screen_P2_PN.md",
                                                       "corpus/spreadsheets/PN_clone_screen_Mar2025.xlsx"], "key_points",
      value=float(mm.group(1)), tol=0.1, key_points=[f"{mm.group(1)} g/L", f"{mm.group(2)} flasks"])
    on, op = C.online("FB-005")
    w = on[(on.t > 2) & (on.t < 40)]
    i = w.pO2.idxmin()
    q("L2", "aggregation", "What was the lowest dissolved-oxygen reading in FB-005 between 2 h and 40 h after inoculation, and when did it occur?",
      f"{w.pO2[i]:.1f} % at about {w.t[i]:.1f} h", [str(op)], "key_points", value=float(w.pO2[i]), tol=0.5,
      key_points=[f"{w.pO2[i]:.1f} %", f"~{w.t[i]:.0f} h"])
    n_may = len(glob.glob(str(C.p("hplc/raw/SEQ-202505*.csv"))))
    assert n_may == tr[tr.Analysed.astype(str).str.startswith("2025-05")].Analysed.nunique()
    q("L2", "aggregation", "How many HPLC sequences were run in May 2025?", str(n_may), ["corpus/hplc/raw/"], "numeric",
      value=n_may, tol=0)
    ch = re.search(r"## Changes from version 1\n- (.+?) added", C.text("sops/SOP-ANA-003_v2.md")).group(1)
    assert ch == "Succinate, furfural and HMF"
    q("L2", "lookup", "Which analytes were added to the HPLC method in version 2 of SOP-ANA-003?",
      ch, ["corpus/sops/SOP-ANA-003_v2.md"], "key_points", key_points=["succinate", "furfural", "HMF"])
    flasks = set()
    wb = openpyxl.load_workbook(C.p("spreadsheets/PN_clone_screen_Mar2025.xlsx"), data_only=True)
    for ws in wb.worksheets[3:]:
        flasks |= {r[1] for r in ws.iter_rows(min_row=3, values_only=True) if r[1]}
    q("L2", "aggregation", "How many shake flasks were in the March clone screen?", str(len(flasks)),
      ["corpus/spreadsheets/PN_clone_screen_Mar2025.xlsx"], "numeric", value=len(flasks), tol=0)

    # ======================= L3: joins across documents ======================
    runs6 = list(RT[RT.strain == "YX-006"].run)
    q("L3", "join", "Which bioreactor runs used strain YX-006, and which clone of YX-006 was it?",
      f"{', '.join(runs6)} ({len(runs6)} runs); clone {c6}", ["corpus/notes/", MD, "corpus/spreadsheets/", REG],
      "key_points", key_points=runs6 + [c6], hops=3)
    p5 = {r: float(md_final(r)) for r in ["FB-012", "FB-013", "FB-014", "FB-015", "FB-016", "FB-017"]}
    rank = sorted(p5, key=p5.get, reverse=True)
    q("L3", "join", "Rank the six C5-fraction fed-batches FB-012 to FB-017 by final xylitol (highest first).",
      " > ".join(f"{r} ({p5[r]:.1f})" for r in rank), [MD], "ranking", value=rank, key_points=rank, hops=2)
    base = [C.ev("FB-001")[1]["Final xylitol (g/L)"], float(md_final("FB-002")), C.ev("FB-003")[1]["Final xylitol (g/L)"]]
    q("L3", "join", "What was the mean and standard deviation of final xylitol in the three baseline runs FB-001 to FB-003?",
      f"{np.mean(base):.1f} +/- {np.std(base, ddof=1):.1f} g/L (FB-001 {base[0]:.1f}, FB-002 {base[1]:.1f}, FB-003 {base[2]:.1f})",
      [ev_path("FB-001"), ev_path("FB-003"), MD, "corpus/meetings/minutes_2025-02-14.md"], "numeric",
      value=float(np.mean(base)), tol=0.3, hops=3)
    counts = RT.operator.value_counts()
    top = counts.index[0]
    q("L3", "join", "Which person operated the most bioreactor runs, and how many?",
      f"{NAME[top]} ({top}), {counts[top]} runs (EV {counts.get('EV', 0)}, PN {counts.get('PN', 0)})",
      ["corpus/notes/", MD], "key_points", key_points=[NAME[top], str(counts[top])], value=int(counts[top]), hops=2)
    og = sorted("FB-" + re.search(r"GA4_fb(\d+)_", Path(f).name).group(1) for f in glob.glob(str(C.p("offgas/*.xlsx"))))
    r12 = sorted(RT[RT.vessel.isin(["R1", "R2"])].run)
    assert og == r12, "off-gas coverage must equal the R1/R2 runs"
    q("L3", "join", "For which bioreactor runs is off-gas (O2/CO2) data available, and what do those runs have in common?",
      f"{', '.join(og)} ({len(og)} runs); they are exactly the runs in vessels R1 and R2 - the gas analyser only "
      "measures the R1/R2 exhaust streams, so no run in R3 or R4 has off-gas data.",
      ["corpus/offgas/", "corpus/notes/", MD], "key_points", key_points=og + ["R1", "R2"], hops=2)
    built = {s_: re.search(r"\*\*Parent:\*\* (\S+)", C.text(f"strains/construction_{s_}.md")).group(1)
             for s_ in ["YX-003", "YX-004"]}
    sel2 = reg.loc["YX-002", "Selected clone"]
    assert built["YX-003"] == built["YX-004"] == "YX-002-c1" and sel2 != "c1"
    q("L3", "provenance", "Which strains were built on a YX-002 clone other than the one finally selected from the YX-002 screen?",
      f"YX-003 and YX-004 were both built on YX-002 c1 (the first PCR-confirmed clone, before the screen); the screen "
      f"selected YX-002 {sel2}, and the team decided to keep c1 as the background.",
      ["corpus/strains/construction_YX-003.md", "corpus/strains/construction_YX-004.md", REG,
       "corpus/meetings/minutes_2025-03-14.md"], "key_points", key_points=["YX-003", "YX-004", "c1", sel2], hops=3)

    def pn_final(sheet):
        rows = C.xlsx_sheet("spreadsheets/PN_FB022-023_lot_comparison.xlsx", sheet)
        data = [r for r in rows[2:] if r[1]]
        return float(data[-1][9])
    a22, a23 = pn_final("FB-022"), pn_final("FB-023")
    d23 = round(a23, 1) - round(a22, 1)
    q("L3", "join", "Compare the final xylitol of FB-022 (no enzyme) and FB-023 (enzyme added). What was the difference?",
      f"FB-022 {a22:.1f} g/L, FB-023 {a23:.1f} g/L; difference {d23:+.1f} g/L (FB-023 higher)",
      ["corpus/spreadsheets/PN_FB022-023_lot_comparison.xlsx", "corpus/notes/FB-023_PN.md"], "numeric",
      value=round(d23, 1), tol=0.2, hops=2)
    r16 = mdr[mdr.Run == "FB16"].iloc[0]
    lot16 = mdm[mdm["For run"] == "FB-016"]["Lot used"].iloc[0]
    assert lot16 == r16.Lot
    q("L3", "join", "Which C5 lot and what C5 fraction were used in the feed of FB-016?", f"{r16.Lot}, {r16.Feed}",
      [MD, "corpus/notes/FB-016_MD.md"], "key_points", key_points=[r16.Lot, "70"], hops=2)
    t4 = C.text("notes/SF_screen_P4_PN.md")
    y5 = {c: float(re.search(rf"\| YX-005 {c}, 50 % C5 \| ([\d.]+) \|", t4).group(1)) for c in ("c1", "c2", "c3")}
    sel5 = reg.loc["YX-005", "Selected clone"]
    assert max(y5, key=y5.get) == sel5
    q("L3", "rationale", "Which YX-005 clone was selected, and on what basis?",
      f"{sel5}: highest xylitol at 72 h in the 50 % C5 flask screen ({', '.join(f'{c} {v:.1f}' for c, v in y5.items())} g/L; "
      "one flask per clone).",
      [REG, "corpus/notes/SF_screen_P4_PN.md", "corpus/spreadsheets/PN_C5_tolerance_Apr2025.xlsx"], "key_points",
      key_points=[sel5, "72 h", "50 %", f"{y5[sel5]:.1f}"], hops=2)
    pn_runs = list(RT[RT.operator == "PN"].run)
    q("L3", "join", "Which bioreactor runs did Priya Nair operate?", ", ".join(pn_runs), ["corpus/notes/"], "key_points",
      key_points=pn_runs, hops=2)

    # ======================= L4: provenance and fault diagnosis ==============
    dev1 = C.text("deviations/DEV-2025-01.md")
    bal, cnt = re.search(r"showed (\d+) mL consumed while the pump counter gave (\d+) mL", dev1).groups()
    pct = re.search(r"delivered about (\d+) %", dev1).group(1)
    mlog = pd.read_csv(C.p("equipment/maintenance_log.csv"), comment="#")
    mpump = mlog[mlog.equipment == "R2 feed pump"].iloc[0]
    q("L4", "fault_diagnosis", "Why was FB-005 excluded from the strain comparison, and what evidence supports it?",
      f"The R2 feed pump over-delivered by about {pct} % (worn tubing): the feed balance showed {bal} mL consumed vs "
      f"{cnt} mL by the pump counter (DEV-2025-01); on {mpump.date} the pump was found to deliver "
      + re.search(r"delivered ([\d.]+ mL/h at [\d.]+ mL/h)", mpump.event).group(1) + " setpoint and was recalibrated. "
      "FB-005 therefore received more glucose and xylose than planned.",
      ["corpus/deviations/DEV-2025-01.md", "corpus/notes/FB-005_EV.md", "corpus/equipment/maintenance_log.csv",
       "corpus/meetings/minutes_2025-04-11.md"], "key_points",
      key_points=["feed pump", f"{pct} %", bal, cnt, "DEV-2025-01"], hops=3)
    on, op = C.online("FB-010")
    dt = on.t.diff()
    i = dt.idxmax()
    dev3 = C.text("deviations/DEV-2025-03.md")
    assert "UPS" in dev3 and "restarted by hand" in dev3
    q("L4", "fault_diagnosis", "What happened in FB-010 around 50 h, and how long is the gap in its online data?",
      f"The lab UPS tripped (power failure): the R4 controller and feed pump were off and the feed was restarted by hand "
      f"(DEV-2025-03, ~3.5 h without feed). The online log has no records between {on.t[i - 1]:.1f} h and {on.t[i]:.1f} h "
      f"(~{dt[i]:.1f} h gap).",
      [str(op), "corpus/notes/FB-010_MD.md", "corpus/deviations/DEV-2025-03.md", "corpus/equipment/maintenance_log.csv"],
      "key_points", value=round(float(dt[i]), 1), tol=0.3, key_points=["power", f"{dt[i]:.1f} h", "feed"], hops=2)
    o9, p9 = C.online("FB-009")
    o6, _ = C.online("FB-006")
    at = lambda o, h: float(o[(o.t > h - 1) & (o.t < h + 1)].pO2.median())
    assert RT.set_index("run").loc["FB-006", "strain"] == RT.set_index("run").loc["FB-009", "strain"]
    mdo = mlog[mlog.equipment == "R3 pO2 probe"].iloc[0]
    q("L4", "fault_diagnosis", "Which P3 run has an unreliable dissolved-oxygen signal in the second half, and what is the evidence?",
      f"FB-009 (R3, YX-003): pO2 falls steadily from about {at(o9, 40):.0f} % at 40 h to {at(o9, 140):.0f} % at 140 h, "
      f"whereas its replicate FB-006 (same strain, same vessel) stays at about {at(o6, 40):.0f}-{at(o6, 140):.0f} % and the "
      f"offline data of FB-009 are normal. The R3 pO2 probe showed slow response and low readings and its membrane was "
      f"replaced on {mdo.date} (maintenance log; minutes 2025-04-11).",
      [str(p9), "corpus/online/FB006_R3.csv", "corpus/equipment/maintenance_log.csv", "corpus/meetings/minutes_2025-04-11.md"],
      "key_points", key_points=["FB-009", "drift", "membrane", mdo.date], hops=3)
    plan24 = mdr[mdr.Run == "FB24"].Lot.iloc[0]
    used24 = mdm[mdm["For run"] == "FB-024"]["Lot used"].iloc[0]
    f24, f22, f25 = (C.tr_value(f"{r}-F", "Formate") for r in ("FB024", "FB022", "FB025"))
    exp02, exp03 = C.coa("C5-2502")["Formic acid"] * 0.6, C.coa("C5-2503")["Formic acid"] * 0.6
    assert abs(f24 - exp02) < abs(f24 - exp03)
    q("L4", "provenance", "Which C5 lot was FB-024 planned with, which lot was actually used, and how can this be shown from the data?",
      f"Planned {plan24} (run sheet and MD note); actually {used24} (MD 'Media & feeds' sheet, DEV-2025-04). The FB-024 feed "
      f"sample contained {f24:.1f} g/L formate - like the C5-2502 feed of FB-022 ({f22:.1f}) and the 60 % of the C5-2502 CoA "
      f"value ({exp02:.1f} g/L), and unlike the C5-2503 feed of FB-025 ({f25:.1f}; CoA-based {exp03:.1f} g/L).",
      [MD, "corpus/hplc/processed/HPLC_results_2025-06.xlsx", "corpus/materials/CoA_C5-2502.md",
       "corpus/materials/CoA_C5-2503.md", "corpus/deviations/DEV-2025-04.md", "corpus/reports/weekly_PN_2025-07-04.md"],
      "key_points", key_points=[plan24, used24, f"{f24:.1f}", "formate"], hops=4)
    rows = tr[(tr.SampleID == "FB005-S07") & (tr.Analyte == "Xylitol")]
    orow = rows[~rows.Comment.fillna("").str.startswith("re-run")].iloc[0]
    crow = rows[rows.Comment.fillna("").str.startswith("re-run")].iloc[0]
    seq = lambda d: f"SEQ-{str(d).replace('-', '')}"
    q("L4", "provenance", "For sample FB005-S07, what xylitol value was first reported, what is the corrected value, and why did it change?",
      f"First {orow.Result_gL:.2f} g/L ({seq(orow.Analysed)}), corrected {crow.Result_gL:.2f} g/L from the re-run "
      f"{seq(crow.Analysed)}: the xylitol standard used in {seq(orow.Analysed)} was prepared from a partially hydrated batch, "
      "so xylitol results of that sequence were biased high (DEV-2025-02).",
      ["corpus/hplc/processed/HPLC_results_2025-03.xlsx", "corpus/hplc/raw/", "corpus/deviations/DEV-2025-02.md"],
      "key_points", key_points=[f"{orow.Result_gL:.2f}", f"{crow.Result_gL:.2f}", "DEV-2025-02"], hops=3)
    bad, stale = [], []
    for _, r in mds.iterrows():
        sid = f"{r.Run.replace('FB', 'FB0')}-{r.Smpl}"
        if isinstance(r.Xol, (int, float)) and not pd.isna(r.Xol):
            first, corr = C.tr_value(sid, "Xylitol", rerun=False), C.tr_value(sid, "Xylitol")
            if first is not None and abs(r.Xol - first) > 0.011:
                bad.append((sid, r.Xol, first))
            elif corr is not None and abs(r.Xol - corr) > 0.011:
                stale.append(sid)
    assert len(bad) == 1
    sid, lv, hp = bad[0]
    lab = lambda x: f"{x[:5].replace('FB0', 'FB-0')} {x[-3:]}"
    st_runs = sorted({lab(x)[:6] for x in stale})
    st_txt = "; ".join(f"{r_} " + ", ".join(lab(x)[-3:] for x in stale if lab(x).startswith(r_)) for r_ in st_runs)
    q("L4", "data_quality", "Is any xylitol value in Mark Doherty's master log inconsistent with the HPLC results? If so, which and why?",
      f"Yes. (1) {lab(sid)} is logged as {lv} g/L, but the HPLC result is {hp:.2f} g/L - a transcription error (decimal "
      f"point shifted, factor 10) that also breaks the time course of the run. (2) {len(stale)} values ({st_txt}) still "
      "show the original results of HPLC sequence SEQ-20250321; they were superseded by the re-run SEQ-20250409 after "
      "DEV-2025-02 (biased xylitol standard), but the master log was not updated.",
      [MD, "corpus/hplc/processed/", "corpus/deviations/DEV-2025-02.md"], "key_points",
      key_points=[lab(sid), f"{lv}", f"{hp:.2f}", "SEQ-20250321", "not updated"] + st_runs, hops=3)
    o14, p14 = C.online("FB-014")
    taf = float(re.search(r"- (\d+)h lots of foam", notes["FB-014_MD"]).group(1))
    pre = o14[(o14.t > taf - 3) & (o14.t < taf)].pO2.mean()
    post = o14[(o14.t > taf + 1) & (o14.t < taf + 4)].pO2.mean()
    q("L4", "fault_diagnosis", "Did the antifoam added to FB-014 affect dissolved oxygen? Quantify.",
      f"Yes: antifoam was added at ~{taf:.0f} h (MD note); pO2 fell from about {pre:.0f} % ({taf - 3:.0f}-{taf:.0f} h) to about "
      f"{post:.0f} % ({taf + 1:.0f}-{taf + 4:.0f} h), consistent with antifoam lowering the oxygen transfer (kLa).",
      ["corpus/notes/FB-014_MD.md", str(p14)], "key_points", value=round(float(pre - post), 0), tol=5,
      key_points=[f"{taf:.0f} h", f"{pre:.0f}", f"{post:.0f}"], hops=2)
    od2 = float(mds[(mds.Run == "FB02") & (mds.Smpl == "S01")].OD.iloc[0])
    od1 = C.ev("FB-001")[0].OD600.iloc[0]
    od3 = C.ev("FB-003")[0].OD600.iloc[0]
    dw2 = float(mds[(mds.Run == "FB02") & (mds.Smpl == "S02")].DW.iloc[0])
    dw1, dw3 = C.ev("FB-001")[0]["DW (g/L)"].iloc[1], C.ev("FB-003")[0]["DW (g/L)"].iloc[1]
    assert "2 days old" in notes["FB-002_MD"]
    q("L4", "rationale", "Why did FB-002 start more slowly than FB-001 and FB-003?",
      f"The preculture was two days old (MD note), giving a lower start OD600 ({od2:.2f} vs {od1:.2f} and {od3:.2f}) and lower "
      f"dry weight at 4 h ({dw2:.2f} vs {dw1:.2f} and {dw3:.2f} g/L); the final titre was nevertheless close to the others.",
      ["corpus/notes/FB-002_MD.md", MD, ev_path("FB-001"), ev_path("FB-003")], "key_points",
      key_points=["preculture", f"{od2:.2f}", f"{dw2:.2f}"], hops=3)
    par = {k: re.search(r"\*\*Parent:\*\* (\S+)", C.text(f"strains/construction_{k}.md")).group(1)
           for k in ("YX-006", "YX-003", "YX-002", "YX-001")}
    q("L4", "provenance", "Trace the lineage of YX-006 back to the parent strain, including the clones used and the genetic change at each step.",
      f"YX-006 = {par['YX-006']} + heterologous FDH (plasmid); YX-003 = {par['YX-003']} + pTEF1-ZWF1 and pTEF1-GND1 "
      f"(oxidative PPP for NADPH); YX-002 = {par['YX-002']} + second, integrated XYL1 copy (pTDH3); YX-001 = the lab host "
      "with plasmid-borne xylose reductase (XYL1).",
      [f"corpus/strains/construction_{k}.md" for k in ("YX-006", "YX-003", "YX-002", "YX-001")] + [REG], "key_points",
      key_points=[par["YX-006"], par["YX-003"], "FDH", "ZWF1", "GND1", "XYL1"], hops=4)

    # ======================= L5: synthesis ===================================
    grp = {"YX-001": ["FB-004", "FB-010"], "YX-002": ["FB-008"], "YX-003": ["FB-006", "FB-009"], "YX-004": ["FB-007", "FB-011"]}
    for k, v in grp.items():
        assert all(RT.set_index("run").loc[r, "strain"] == k for r in v)
    assert RT.set_index("run").loc["FB-005", "strain"] == "YX-002"
    fin = {r: C.final(r) for v in grp.values() for r in v}
    mean = {k: np.mean([fin[r] for r in v]) for k, v in grp.items()}
    gains = {k: 100 * (mean[k] / mean["YX-001"] - 1) for k in grp}
    best = max([k for k in grp if k != "YX-001"], key=mean.get)
    sd1 = float(np.std(base, ddof=1))
    ctrl_wo10 = fin["FB-004"]
    g_wo10 = 100 * (mean[best] / ctrl_wo10 - 1)
    q("L5", "synthesis", "In the P3 bioreactor comparison (FB-004 to FB-011, excluding FB-005, using corrected HPLC values), "
      "which engineered strain gave the highest mean final xylitol, by how much compared with the parent control, and how "
      "reliable is that difference?",
      f"{best}: {mean[best]:.1f} g/L vs control {mean['YX-001']:.1f} g/L ({gains[best]:+.1f} %); YX-003 {mean['YX-003']:.1f}, "
      f"YX-002 {mean['YX-002']:.1f} (FB-008 only, because FB-005 is excluded). The evidence is weak: only 1-2 runs per "
      f"strain, a difference of {mean[best] - mean['YX-001']:.1f} g/L against a run-to-run sd of about {sd1:.1f} g/L "
      f"(baseline FB-001 to FB-003), and the control mean includes FB-010 ({fin['FB-010']:.1f} g/L), which had ~3.5 h without "
      f"feed in the power cut (DEV-2025-03). Against FB-004 alone ({ctrl_wo10:.1f} g/L) the gain shrinks to "
      f"{g_wo10:+.1f} %. Round-1 values also had to be corrected for the HPLC standard error (DEV-2025-02).",
      ["corpus/hplc/processed/", "corpus/lims/sample_register.csv", "corpus/deviations/DEV-2025-01.md",
       "corpus/deviations/DEV-2025-02.md", "corpus/deviations/DEV-2025-03.md", "corpus/meetings/minutes_2025-02-14.md",
       "corpus/meetings/minutes_2025-04-11.md"],
      "key_points", value=round(gains[best], 1), tol=1.0,
      key_points=[best, f"{mean[best]:.1f}", f"{mean['YX-001']:.1f}", "FB-010", "few runs"], hops=5)

    m05 = re.search(r"At 70 % C5 the parent YX-001 made (no detectable xylitol|[\d.]+ g/L xylitol)(?: and did not grow)? "
                    r"\(OD ([\d.]+) at 72 h(?:, inoculated at OD ([\d.]+))?\)", minutes["2025-05-02"])
    x12 = float(md_final("FB-012")); d12 = float(md_final("FB-012", "DW")); f12 = float(md_final("FB-012", "For"))
    stall = re.search(r"stopped working by ~(\d+)h", notes["FB-012_MD"]).group(1)
    fdh70 = {r: float(md_final(r)) for r in ("FB-015", "FB-016", "FB-017")}
    same_bg = [r for r in fdh70 if RT.set_index("run").loc[r, "strain"] == "YX-006"]
    q("L5", "synthesis", "Is the heterologous FDH needed for xylitol production at 70 % C5 fraction? Support the answer with flask and bioreactor data.",
      f"Yes. In flasks at 70 % C5 the parent YX-001 made {m05.group(1)} and did not grow (OD {m05.group(2)} at 72 h"
      + (f" from OD {m05.group(3)}" if m05.group(3) else "") + "), while the FDH strains grew and made xylitol. In the "
      f"bioreactor, FB-012 (YX-003, no FDH, 70 % C5) stopped working by ~{stall} h (xylose accumulating, formate "
      f"{f12:.1f} g/L at the end) and finished at only {x12:.1f} g/L xylitol, DW {d12:.1f} g/L; with FDH at 70 % C5 the runs "
      f"reached {', '.join(f'{k} {v:.1f}' for k, v in fdh70.items())} g/L - {' and '.join(same_bg)} with YX-006, which is "
      "YX-003 plus FDH, i.e. the same background as FB-012.",
      ["corpus/spreadsheets/PN_C5_tolerance_Apr2025.xlsx", "corpus/meetings/minutes_2025-05-02.md", MD,
       "corpus/notes/FB-012_MD.md", "corpus/strains/construction_YX-006.md", "corpus/meetings/minutes_2025-05-30.md"],
      "key_points", key_points=["yes", "formate", "FB-012", f"{x12:.1f}"] + same_bg[:1], hops=5)

    p6 = {r: C.ev(r) for r in ("FB-018", "FB-019", "FB-020", "FB-021")}
    p6v = {r: v[1]["Final xylitol (g/L)"] for r, v in p6.items()}
    p6c = {r: v[2]["Purpose"] for r, v in p6.items()}
    assert "500" in p6c["FB-021"] and "B" in p6c["FB-021"] and "900" in p6c["FB-020"]
    p6o = {r: float(C.online(r)[0].query("t > 30").pO2.median()) for r in p6v}
    def fed(r):
        d = p6[r][0]
        d = d[d.Sample.str.contains("-S")].copy()
        d["Time (h)"] = d["Time (h)"].astype(float)
        return d[d["Time (h)"] > 40]
    e21 = fed("FB-021")
    late = e21[e21["Ethanol (g/L)"].fillna(0) > 0.5]
    others_late = {r: float(fed(r)["Ethanol (g/L)"].fillna(0).max()) for r in ("FB-018", "FB-019", "FB-020")}
    assert len(late) and max(others_late.values()) < 0.5
    q("L5", "synthesis", "In the P6 study with YX-006 (900 vs 500 rpm, feed A vs the more glucose-rich feed B), which "
      "condition was oxygen-limited, what is the evidence, and did it measurably reduce the final xylitol titre?",
      f"FB-021 (500 rpm, feed B): its fed-phase pO2 median was {p6o['FB-021']:.0f} % (vs {p6o['FB-018']:.0f} / "
      f"{p6o['FB-019']:.0f} / {p6o['FB-020']:.0f} % for FB-018/019/020), and ethanol re-appeared from "
      f"{late['Time (h)'].iloc[0]:g} h ({late['Ethanol (g/L)'].iloc[-1]:.1f} g/L at the end) while it stayed below detection "
      f"in the other three runs - overflow/fermentative metabolism under O2 limitation. Final xylitol was "
      f"{p6v['FB-018']:.1f} / {p6v['FB-019']:.1f} / {p6v['FB-020']:.1f} / {p6v['FB-021']:.1f} g/L; the "
      f"{p6v['FB-020'] - p6v['FB-021']:.1f} g/L lower titre of FB-021 vs FB-020 is a single-run comparison within the "
      f"run-to-run scatter (baseline sd {sd1:.1f} g/L), so a titre loss is not demonstrated.",
      [ev_path(r) for r in p6] + ["corpus/notes/FB-021_EV.md", "corpus/online/", "corpus/meetings/minutes_2025-06-20.md",
                                  "corpus/meetings/minutes_2025-02-14.md"], "key_points",
      key_points=["FB-021", f"{p6o['FB-021']:.0f} %", "ethanol", "not significant"], hops=5)

    f25 = md_final("FB-025"); f24x = md_final("FB-024")
    q("L5", "synthesis", "Taking the lot mix-up into account, what can be concluded about the effect of the enzyme on lot C5-2503?",
      f"Nothing: the only real C5-2503 run is FB-025 (with enzyme, {f25} g/L); FB-024, meant as the no-enzyme C5-2503 "
      f"control, was fed with C5-2502 ({f24x} g/L), so there is no C5-2503 run without enzyme to compare against "
      f"(a repeat was planned for August). For C5-2502 the enzyme changed final xylitol by {d23:+.1f} g/L (FB-023 vs FB-022, "
      "single runs).",
      ["corpus/deviations/DEV-2025-04.md", MD, "corpus/spreadsheets/PN_FB022-023_lot_comparison.xlsx",
       "corpus/meetings/minutes_2025-07-11.md"], "key_points",
      key_points=["cannot", "FB-025", "FB-024", "C5-2502"], hops=4)

    def ser(run, col):
        return dict(zip(mds[mds.Run == run].Smpl, mds[mds.Run == run][col]))
    x12s, x13s, f12s, f13s = ser("FB12", "Xol"), ser("FB13", "Xol"), ser("FB12", "For"), ser("FB13", "For")
    lims = C.lims().set_index("sample_id")
    t12, t13 = lims.loc["FB012-S07", "sampled_at"], lims.loc["FB013-S07", "sampled_at"]
    num = lambda v: isinstance(v, (int, float))
    f13_other = max([f13s[k] for k in f13s if k != "S07" and num(f13s[k])] or [0.0])
    assert not num(f12s["S07"]) and num(f13s["S07"]) and f13s["S07"] > 20 * max(f13_other, 0.05)
    q("L5", "data_quality", "Do any HPLC results of FB-012 or FB-013 break the time course in a way that suggests a sample mix-up? Explain.",
      f"Yes, sample S07 (72 h) of both runs. FB-012 S07 shows xylitol {x12s['S07']} g/L (S06 {x12s['S06']}, S08 {x12s['S08']}) "
      f"and no formate although formate is present in all its other fed-phase samples; FB-013 S07 shows xylitol "
      f"{x13s['S07']} g/L (S06 {x13s['S06']}, S08 {x13s['S08']}) and formate {f13s['S07']} g/L although all other FB-013 "
      f"samples have at most {f13_other:.2f} g/L formate. Each S07 looks like the other run's sample, so the two tubes were most likely swapped "
      f"before HPLC (both sampled the same morning, {t12} and {t13}). OD and dry weight of S07 fit their own runs.",
      [MD, "corpus/hplc/processed/HPLC_results_2025-05.xlsx", "corpus/lims/sample_register.csv"], "key_points",
      key_points=["S07", "swap", "FB-012", "FB-013", "formate"], hops=3)

    finals = {r: C.final(r) for r in RT.run}
    order = sorted(finals, key=finals.get, reverse=True)
    top1 = order[0]
    if top1 == "FB-005":
        nxt = order[1]
        ans = (f"FB-005 ({finals['FB-005']:.1f} g/L), but it is not trustworthy: its feed pump over-delivered by ~15 % "
               f"(DEV-2025-01), so it received more substrate and was excluded. The highest valid run is {nxt} "
               f"({finals[nxt]:.1f} g/L), followed by " + ", ".join(f"{r} ({finals[r]:.1f})" for r in order[2:4])
               + " - differences between these are within run-to-run scatter.")
        kp = ["FB-005", f"{finals['FB-005']:.1f}", "DEV-2025-01", nxt]
    else:
        ans = (f"{top1} ({finals[top1]:.1f} g/L); no deviation is recorded for it, but it is a single run and the next runs "
               f"({', '.join(f'{r} {finals[r]:.1f}' for r in order[1:3])}) are within run-to-run scatter.")
        kp = [top1, f"{finals[top1]:.1f}"]
    q("L5", "synthesis", "Which bioreactor run reached the highest final xylitol in the whole campaign, and is that result trustworthy?",
      ans, ["corpus/hplc/processed/", "corpus/lims/sample_register.csv", "corpus/deviations/", "corpus/spreadsheets/"],
      "key_points", key_points=kp, hops=4)

    # ======================= U: unanswerable / false premise =================
    q("U", "unanswerable", "What was the dissolved CO2 concentration in the broth of FB-019?",
      "Not available: dissolved CO2 was not measured. Only the exhaust-gas CO2 (%) was recorded for FB-019 (off-gas file).",
      ["corpus/offgas/GA4_fb019_R2.xlsx", "corpus/sops/"], "unanswerable", key_points=["not measured"])
    t6 = tr[tr.Run == "FB-006"]
    assert set(t6.SOP) == {"v1"} and "Succinate" not in set(t6.Analyte)
    q("U", "unanswerable", "What succinate concentration was measured in FB-006?",
      "Not available: FB-006 (March 2025) samples were analysed with SOP-ANA-003 v1, which does not include succinate; "
      "succinate was only added in v2, effective 2025-06-01.", ["corpus/sops/SOP-ANA-003_v1.md", "corpus/sops/SOP-ANA-003_v2.md",
                                                                 "corpus/hplc/processed/HPLC_results_2025-03.xlsx"],
      "unanswerable", key_points=["not measured", "v1", "2025-06-01"])
    assert RT.set_index("run").loc["FB-006", "vessel"] == "R3" and "FB-006" not in og
    q("U", "unanswerable", "What was the respiratory quotient (RQ) of FB-006 during the fed phase?",
      "Not available: there is no off-gas data for FB-006 (vessel R3); off-gas was only recorded for R1 and R2 runs.",
      ["corpus/offgas/"], "unanswerable", key_points=["no off-gas", "R3"])
    q("U", "unanswerable", "What final xylitol titre did run FB-026 reach?",
      f"There is no run FB-026 in the records (runs go from FB-001 to {RT.run.max()}); the minutes of 2025-07-11 only plan "
      "a repeat for August.", ["corpus/lims/sample_register.csv", "corpus/meetings/minutes_2025-07-11.md"],
      "unanswerable", key_points=["no such run"])
    stat4 = reg.loc["YX-004", "Status"]
    if not str(stat4).startswith("dropped"):
        q("U", "false_premise", "Why was strain YX-004 dropped after the P3 comparison?",
          f"False premise: YX-004 was not dropped. The registry lists it as '{stat4}', and the 2025-04-11 minutes decided to "
          "keep it on hold.", [REG, "corpus/meetings/minutes_2025-04-11.md"], "unanswerable",
          key_points=["not dropped", stat4])
    else:
        q("U", "false_premise", "Why was strain YX-003 dropped after the P3 comparison?",
          "False premise: YX-003 was not dropped; it became the background of YX-006.", [REG], "unanswerable",
          key_points=["not dropped"])
    q("U", "unanswerable", "How much did lot C5-2502 of the C5 fraction cost?",
      "Not available: no document gives prices or costs for the C5 lots.", ["corpus/materials/"], "unanswerable",
      key_points=["not documented"])
    q("U", "false_premise", "What final xylitol did the C5-2503 run without enzyme reach?",
      "There is no such run: FB-024 was planned as the no-enzyme C5-2503 run but was fed with C5-2502 (DEV-2025-04). "
      "The only C5-2503 run, FB-025, had enzyme.", ["corpus/deviations/DEV-2025-04.md", MD], "unanswerable",
      key_points=["no such run", "FB-024", "C5-2502"])
    return Q


def finalize(Q, C):
    """Core questions get set/interop fields; then the IO / CA / XD sets are appended."""
    from questions_extra import CORE_TAGS, build_extra
    for x in Q:
        x.setdefault("set", "core")
        x["interop"] = CORE_TAGS.get(x["id"], [])
        x.setdefault("trap", None)
        x.setdefault("rubric", None)
        x.setdefault("penalties", None)
    build_extra(C, Q)
    return Q


SET_NAMES = {"core": "Core set (difficulty tiers L1-L5, U)", "IO": "Semantic interoperability (IO)",
             "CA": "Calculations (CA)", "XD": "Experiment design (XD)"}


def write(Q, out: Path):
    from questions_extra import INTEROP_TAGS
    out.mkdir(parents=True, exist_ok=True)
    (out / "questions.json").write_text(json.dumps(Q, indent=1, default=str))
    flat = lambda v: "; ".join(map(str, v)) if isinstance(v, list) else v
    pd.DataFrame([{**x, "evidence": flat(x["evidence"]), "key_points": flat(x["key_points"]), "interop": flat(x["interop"]),
                   "rubric": json.dumps(x["rubric"]) if x["rubric"] else None,
                   "penalties": flat(x["penalties"]) if x["penalties"] else None} for x in Q]).to_csv(out / "questions.csv", index=False)
    md = ["# Pilot benchmark questions\n",
          "Gold answers are computed from the corpus files by `generator/questions.py` and `generator/questions_extra.py`. "
          "Difficulty tiers: L1 lookup, L2 single document + calculation, L3 joins, L4 provenance / faults / harder "
          "calculation, L5 synthesis, U unanswerable or false premise.\n",
          "Interoperability tags: " + "; ".join(f"`{k}` - {v}" for k, v in INTEROP_TAGS.items()) + ".\n"]
    for st in ("core", "IO", "CA", "XD"):
        md.append(f"\n## {SET_NAMES[st]}\n")
        for x in [y for y in Q if y["set"] == st]:
            md.append(f"\n### {x['id']} ({x['tier']}, {x['category']}, {x['hops']} hop{'s' if x['hops'] > 1 else ''})\n\n"
                      f"**Q:** {x['question']}\n\n**A:** {x['gold_answer']}\n")
            if x["rubric"]:
                md.append("\n| Item | Suggestion | Evidence in the corpus | Simulator check (answer key) |\n|---|---|---|---|\n"
                          + "".join(f"| {r['id']} | {r['suggestion']} | {r['evidence']} | {r['simulator_check']} |\n" for r in x["rubric"]))
                md.append("\nPenalised: " + " ".join(f"({i + 1}) {p}" for i, p in enumerate(x["penalties"])) + "\n")
            md.append(f"\n_Evidence:_ {', '.join(x['evidence'])}  \n_Grading:_ {x['grading']}"
                      + (f" (key points: {', '.join(map(str, x['key_points']))})" if x["key_points"] else "")
                      + (f"  \n_Interop:_ {', '.join(x['interop'])}" if x["interop"] else "")
                      + (f"  \n_Trap:_ {x['trap']}" if x["trap"] else "") + "\n")
    (out / "questions.md").write_text("\n".join(md))


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--pilot", default=str(Path(__file__).resolve().parents[1] / "pilot"))
    pilot = Path(ap.parse_args().pilot)
    C = Corpus(pilot / "corpus")
    Q = finalize(build_questions(C), C)
    write(Q, pilot / "questions")
    print(pd.Series([f"{x['set']}:{x['tier']}" for x in Q]).value_counts().sort_index().to_dict(), "total", len(Q))
