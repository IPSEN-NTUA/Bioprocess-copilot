"""
Renderers for instrument exports and spreadsheets. Every file starts with the synthetic-data notice.
"""
from __future__ import annotations

import datetime as dtm
from pathlib import Path

import numpy as np
import openpyxl
from openpyxl.styles import Font

import measure as ms
import world as W

LONG_LEAD = {"FB-012": 17.4, "FB-021": 15.8}       # loggers started the day before inoculation
ANALYTE_ORDER = ["Glc", "Xyl", "Ara", "Xol", "Gly", "Ac", "EtOH", "For", "Pyr", "Suc", "Fur", "HMF"]


def _fmt(v, nd=2):
    return "" if v is None else f"{v:.{nd}f}"


def _new_wb():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "About"
    ws["A1"] = W.SYNTHETIC_NOTICE
    ws["A1"].font = Font(bold=True)
    return wb


# ---------------------------------------------------------------------------
# Online logs (two vendors)
# ---------------------------------------------------------------------------
def online_bioctrl(rec, path: Path, rng):
    spec, on = rec["spec"], rec["online"]
    lead = rec.setdefault("logger_lead_h", float(rng.uniform(0.2, 1.5)))
    if spec.run_id in LONG_LEAD:                          # logger left running from the previous day
        lead = rec["logger_lead_h"] = LONG_LEAD[spec.run_id]
    start = spec.start - dtm.timedelta(hours=lead)
    lines = [f"# {W.SYNTHETIC_NOTICE}",
             f"Instrument\t: BioCtrl B1\t\tVessel\t: {spec.vessel}",
             f"Batch name\t: {spec.run_id}\t\tOperator\t: {spec.operator}",
             f"Log start\t: {start:%d %b %Y %H:%M:%S}",
             f"Inoculation\t: {spec.start:%d %b %Y %H:%M:%S}",
             "Date\tTime\tStirrer[rpm]\tTemp[C]\tpH\tpO2[%]\tAcid_Pump[cnt]\tBase_Pump[cnt]\tFeed_Pump[cnt]"]
    t = -lead
    while t < 0:                                           # sterile medium before inoculation
        ts = spec.start + dtm.timedelta(hours=t)
        lines.append(f"{ts:%d.%m.%Y}\t{ts:%H:%M:%S}\t{300 + rng.normal(0, 0.5):.0f}\t{30 + rng.normal(0, 0.02):.2f}\t"
                     f"{on['pH'].iloc[0]:.3f}\t{100 + rng.normal(0, 0.3):.1f}\t0\t0\t0")
        t += 1 / 6
    for _, r in on.iterrows():
        ts = spec.start + dtm.timedelta(hours=float(r.t))
        lines.append(f"{ts:%d.%m.%Y}\t{ts:%H:%M:%S}\t{r.stirrer:.0f}\t{r.temp:.2f}\t{r.pH:.3f}\t{r.pO2:.1f}\t"
                     f"{int(r.acid_counts)}\t{int(r.base_counts)}\t{int(r.feed_counts)}")
    path.write_text("\r\n".join(lines) + "\r\n", encoding="utf-8")


def online_vesselview(rec, path: Path, rng):
    spec, on = rec["spec"], rec["online"]
    lead = rec.setdefault("logger_lead_h", float(rng.uniform(0.2, 1.5)))
    if spec.run_id in LONG_LEAD:
        lead = rec["logger_lead_h"] = LONG_LEAD[spec.run_id]
    rows = [f"# {W.SYNTHETIC_NOTICE}",
            f"# VesselView VV-1 export | unit {spec.vessel} | recipe {spec.run_id} | exported by {spec.operator}",
            "Timestamp,Elapsed_h,AGIT_rpm,TEMP_C,PH,DO_pct,FEED_BAL_g,BASE_mL,ACID_mL"]
    t = -lead
    bal0 = on["feed_balance_g"].iloc[0]
    while t < 0:
        ts = spec.start + dtm.timedelta(hours=t)
        rows.append(f"{ts:%Y-%m-%dT%H:%M:%S},{t + lead:.3f},{300 + rng.normal(0, 0.5):.0f},{30 + rng.normal(0, 0.02):.2f},"
                    f"{on['pH'].iloc[0]:.2f},{100 + rng.normal(0, 0.3):.1f},{bal0 + rng.normal(0, 0.05):.2f},0.00,0.00")
        t += 1 / 6
    for _, r in on.iterrows():
        ts = spec.start + dtm.timedelta(hours=float(r.t))
        rows.append(f"{ts:%Y-%m-%dT%H:%M:%S},{r.t + lead:.3f},{r.stirrer:.0f},{r.temp:.2f},{r.pH:.2f},{r.pO2:.1f},"
                    f"{r.feed_balance_g:.2f},{r.base_counts * ms.PUMP_ML_PER_COUNT['base']:.2f},"
                    f"{r.acid_counts * ms.PUMP_ML_PER_COUNT['acid']:.2f}")
    path.write_text("\n".join(rows) + "\n", encoding="utf-8")


# ---------------------------------------------------------------------------
# Off-gas analyser export (xlsx)
# ---------------------------------------------------------------------------
def offgas_xlsx(rec, path: Path):
    spec, og = rec["spec"], rec["offgas"]
    wb = _new_wb()
    ws = wb.create_sheet("GA4")
    t_shift = og["rel_time_s"].iloc[0] / 3600 - og["t"].iloc[0]
    t_start = spec.start - dtm.timedelta(hours=t_shift)
    ws.append(["GasLab GA-4", "export", "", "stream", spec.vessel])
    ws.append(["Run tag", spec.run_id.lower().replace("-", ""), "", "airflow (L/min)", 0.7])
    ws.append(["Analyser start", f"{t_start:%Y-%m-%d %H:%M}", "", "Inlet air", "O2 20.95 %, CO2 0.04 %"])
    ws.append([])
    ws.append(["Cycle", "Date", "Time", "RelTime[s]", "N2 [%]", "O2 [%]", "Ar [%]", "CO2 [%]"])
    for i, r in og.iterrows():
        ts = t_start + dtm.timedelta(seconds=float(r.rel_time_s))
        ws.append([i + 1, ts.strftime("%Y-%m-%d"), ts.strftime("%H:%M:%S"), r.rel_time_s, r.N2, r.O2, r.Ar, r.CO2])
    wb.save(path)


# ---------------------------------------------------------------------------
# HPLC: raw sequence CSVs and TR's tidy processed workbook
# ---------------------------------------------------------------------------
def hplc_sequence_csv(seq_date, seq, path: Path):
    analytes = ms.HPLC_ANALYTES_V2 if seq["version"] == 2 else ms.HPLC_ANALYTES_V1
    head = [f"# {W.SYNTHETIC_NOTICE}",
            f"# {W.HPLC_ID if seq['version'] == 2 else W.HPLC_ID.replace(' + UV', '')} | sequence SEQ-{seq_date:%Y%m%d} | method SOP-ANA-003 v{seq['version']}"
            + (f" | re-analysis of SEQ-{seq['rerun_of']:%Y%m%d} with fresh standards" if seq.get("rerun_of") else ""),
            "Inj,Vial,SampleID,Dilution,Analyte,RT_min,Area,Conc_vial_gL,Flag"]
    rows, inj = [], 1
    for j, s in enumerate(seq["injections"]):
        for a in analytes:
            r = s["results"][a]
            rows.append(f"{inj},{j + 1},{s['sample_id']},{s['dilution']},{ms.HPLC_NAMES[a]},{ms.RT[a]:.2f},"
                        f"{r['area']},{max(r['vial'], 0):.4f},{r['flag']}")
        inj += 1
    path.write_text("\n".join(head + rows) + "\n", encoding="utf-8")


def hplc_processed_xlsx(samples, month: str, path: Path, sequences):
    wb = _new_wb()
    ws = wb.create_sheet("Results")
    ws.append(["SampleID", "Run", "Sampled", "Analysed", "SOP", "Dilution", "Analyte", "Result_gL", "Flag", "Comment"])
    for s in samples:
        analytes = ms.HPLC_ANALYTES_V2 if s["sop_version"] == 2 else ms.HPLC_ANALYTES_V1
        for a in analytes:
            v = s["hplc"].get(a)
            comment = ""
            if s.get("hplc_rerun") and a in s["hplc_rerun"]:
                comment = f"superseded by re-run {s['hplc_rerun_date']:%Y-%m-%d}" if a == "Xol" else ""
            ws.append([s["sample_id"], s["run_id"], s["datetime"].strftime("%Y-%m-%d %H:%M"),
                       s["hplc_date"].strftime("%Y-%m-%d"), f"v{s['sop_version']}", s["dilution"],
                       ms.HPLC_NAMES[a], None if v is None else round(v, 3), "<LOD" if v is None else "", comment])
        if s.get("hplc_rerun"):
            v = s["hplc_rerun"].get("Xol")
            ws.append([s["sample_id"], s["run_id"], s["datetime"].strftime("%Y-%m-%d %H:%M"),
                       s["hplc_rerun_date"].strftime("%Y-%m-%d"), "v1", s["dilution"], "Xylitol",
                       None if v is None else round(v, 3), "", "re-run, new standard lot"])
    wb.save(path)


# ---------------------------------------------------------------------------
# LIMS sample register (CSV)
# ---------------------------------------------------------------------------
def sample_register_csv(campaign, path: Path):
    rows = [f"# {W.SYNTHETIC_NOTICE}",
            "sample_id,run_id,vessel,sampled_at,process_time_h,type,requested,volume_mL,stored,comment"]
    for rec in campaign.runs:
        spec = rec["spec"]
        for s in rec["samples"]:
            req = "HPLC" if s["kind"] != "broth" else ("HPLC;DW;OD" if spec.kind == "reactor" else "HPLC;OD")
            vol = "2" if s["kind"] != "broth" else ("7" if spec.kind == "reactor" else "1")
            t = "" if s["kind"] != "broth" else f"{s['t']:.1f}"
            rows.append(f"{s['sample_id']},{spec.run_id},{spec.vessel},{s['datetime']:%Y-%m-%d %H:%M},{t},{s['kind']},"
                        f"{req},{vol},-20C box {spec.phase},")
    path.write_text("\n".join(rows) + "\n", encoding="utf-8")


# ---------------------------------------------------------------------------
# Person-specific spreadsheets
# ---------------------------------------------------------------------------
def _hval(s, a, rerun=False):
    src = s.get("hplc_rerun") if rerun and s.get("hplc_rerun") else s["hplc"]
    return src.get(a)


def ev_run_workbook(rec, path: Path, summary):
    """EV: one tidy workbook per run, units in headers, calculated sheet."""
    spec = rec["spec"]
    wb = _new_wb()
    info = wb.create_sheet("Run info")
    strain = f"{spec.strain_id} (clone c{rec['clone'] + 1})" if rec["clone"] is not None else spec.strain_id
    for k, v in [("Run", spec.run_id), ("Vessel", f"{spec.vessel} - {W.VESSELS[spec.vessel].model}"),
                 ("Operator", W.PEOPLE[spec.operator].name), ("Strain", strain), ("Purpose", spec.purpose),
                 ("Inoculation", spec.start.strftime("%Y-%m-%d %H:%M")), ("Batch medium", f"{W.YP} + glucose 20 g/L"),
                 ("Feed", ", ".join(f"{k} {v:.0f} g/L" for k, v in spec.feed.items() if k != "Trc") +
                  (f" ({int(spec.frac*100)} % C5 lot {spec.lot})" if spec.lot else "")),
                 ("Feed start / rate", f"{spec.feed_start_h:.1f} h / {spec.feed_rate_mLh:.1f} mL/h (setpoint)"),
                 ("T / pH / agitation / air", f"30 C / {spec.pH} / {spec.stirrer:.0f} rpm / 0.7 L/min"),
                 ("Initial volume", "300 mL")] + ([("Note", "Xylitol for " + ", ".join(
                     s["sample_id"] for s in rec["samples"] if s.get("hplc_rerun")) +
                     " taken from the HPLC re-run after DEV-2025-02")] if any(s.get("hplc_rerun") for s in rec["samples"]) else []):
        info.append([k, v])
    off = wb.create_sheet("Offline data")
    an = [a for a in ANALYTE_ORDER if any(s["hplc"].get(a) is not None for s in rec["samples"])]
    off.append(["Sample", "Time (h)", "Date", "Volume (mL)", "OD600", "DW (g/L)", "DW sem (g/L)"] +
               [f"{ms.HPLC_NAMES[a]} (g/L)" for a in an])
    for s in rec["samples"]:
        if s["kind"] != "broth":
            continue
        off.append([s["sample_id"], s["t"], s["datetime"].strftime("%Y-%m-%d %H:%M"), round(s["V_mL"] / 10) * 10,
                    s["od"]["od"], round(s["dw"]["mean"], 2), round(s["dw"]["sem"], 3)] +
                   [None if _hval(s, a, rerun=True) is None else round(_hval(s, a, rerun=True), 2) for a in an])
    for s in rec["samples"]:
        if s["kind"] in ("medium", "feed"):
            off.append([s["sample_id"], s["kind"], "", "", "", "", ""] +
                       [None if s["hplc"].get(a) is None else round(s["hplc"][a], 2) for a in an])
    calc = wb.create_sheet("Calculated")
    for k, v in summary.items():
        calc.append([k, v])
    wb.save(path)


def md_master_log(campaign, path: Path, summaries):
    """MD: one running workbook, appended rows, terse headers, text values, a transcription slip."""
    wb = _new_wb()
    runs = wb.create_sheet("Runs")
    runs.append(["Run", "Vsl", "Strain", "Start", "Medium", "Feed", "Lot", "rpm", "Notes"])
    smp = wb.create_sheet("Samples")
    smp.append(["Run", "Smpl", "Date", "h", "OD", "DW", "Glc", "Xyl", "Xol", "Gly", "EtOH", "Ac", "For", "Comment"])
    media = wb.create_sheet("Media & feeds")
    media.append(["Prepared", "Batch", "What", "Lot used", "Vol (mL)", "For run", "By"])
    mine = [r for r in campaign.runs if r["spec"].operator == "MD"]
    for rec in mine:
        spec = rec["spec"]
        feed = "A" if spec.feed == {"Glc": 100.0, "Xyl": 300.0} else ("B" if spec.feed.get("Glc") == 150 else f"C5 {int(spec.frac*100)}%")
        runs.append([spec.run_id.replace("FB-0", "FB"), spec.vessel, spec.strain_id, spec.start.strftime(W.PEOPLE["MD"].date_fmt),
                     "YP+20 glc", feed, spec.lot or "", int(spec.stirrer), summaries[spec.run_id].get("md_note", "")])
        prep = spec.start - dtm.timedelta(days=1)
        media.append([prep.strftime("%d/%m/%y"), f"F{prep:%m%d}-{spec.vessel}", "feed " + feed,
                      spec.lot_actual or "-", 500, spec.run_id, "MD"])
        for s in rec["samples"]:
            if s["kind"] != "broth":
                continue
            def h(a):
                v = s["hplc"].get(a)
                return "<0.05" if v is None else round(v, 2)
            comment = ""
            row = [spec.run_id.replace("FB-0", "FB"), s["sample_id"].split("-")[-1], s["datetime"].strftime("%d/%m/%y"),
                   round(s["t"], 1), s["od"]["od"], round(s["dw"]["mean"], 2), h("Glc"), h("Xyl"), h("Xol"), h("Gly"),
                   h("EtOH"), h("Ac"), h("For"), comment]
            if spec.run_id == "FB-015" and s["sample_id"].endswith("S08") and isinstance(row[8], float):
                row[8] = round(row[8] * 10, 1)                 # transcription slip: decimal point lost
            smp.append(row)
    wb.save(path)


def pn_wide_workbook(runs, title: str, description: list, methods: list, path: Path, flask: bool):
    """PN: wide sheets; several vessels stacked per sheet, notes in column A."""
    wb = _new_wb()
    d = wb.create_sheet("Description")
    for line in description:
        d.append([line])
    if methods:
        m = wb.create_sheet("Methods")
        for row in methods:
            m.append(row if isinstance(row, list) else [row])
    groups = {}
    for rec in runs:
        key = rec["spec"].strain_id if flask else rec["spec"].run_id
        groups.setdefault(key, []).append(rec)
    for key, recs in groups.items():
        ws = wb.create_sheet(key[:31])
        if flask:
            ws.append(["", "", "", "Biomass", "HPLC data (g/L)"])
            ws.append(["Note", "Flask", "Time (h)", "OD600", "Glucose", "Xylose", "Xylitol", "Glycerol", "Ethanol",
                       "Acetate", "Formate"])
        else:
            ws.append(["", "", "", "Feed", "", "Biomass", "", "HPLC data (g/L)"])
            ws.append(["Note", "Sample name", "Time (h)", "Feed added (mL)", "Reactor volume (mL)", "DW (g/L)", "sem",
                       "Glucose", "Xylose", "Xylitol", "Glycerol", "Ethanol", "Acetate", "Formate"])
        for rec in recs:
            spec = rec["spec"]
            notes = [f"{spec.strain_id} c{rec['clone'] + 1}" if rec["clone"] is not None else spec.strain_id,
                     spec.purpose]
            if spec.lot:
                notes.append(f"C5 lot {spec.lot}")
            if not spec.buffered:
                notes.append("no CaCO3")
            bro = [s for s in rec["samples"] if s["kind"] == "broth"]
            for i, s in enumerate(bro):
                h = lambda a: None if s["hplc"].get(a) is None else round(s["hplc"][a], 3)
                note = notes[i] if i < len(notes) else None
                if flask:
                    ws.append([note, spec.run_id, s["t"], s["od"]["od"], h("Glc"), h("Xyl"), h("Xol"), h("Gly"),
                               h("EtOH"), h("Ac"), h("For")])
                else:
                    on = rec.get("online")
                    fed = None
                    if on is not None:
                        fed = float(np.interp(s["t"], on["t"], on["feed_counts"])) * ms.PUMP_ML_PER_COUNT["feed"]
                    ws.append([note, s["sample_id"], s["t"], None if fed is None else round(fed, 1),
                               round(s["V_mL"] / 10) * 10, round(s["dw"]["mean"], 2), round(s["dw"]["sem"], 3),
                               h("Glc"), h("Xyl"), h("Xol"), h("Gly"), h("EtOH"), h("Ac"), h("For")])
            ws.append([])
    wb.save(path)


def strain_registry_xlsx(campaign, path: Path, status: dict):
    wb = _new_wb()
    ws = wb.create_sheet("Strains")
    ws.append(["Strain ID", "Alias", "Parent", "Genotype / change", "Constructed", "By", "Clones tested",
               "Selected clone", "Cryo position", "Status", "Notes"])
    aliases = {"YX-001": "parent", "YX-002": "2xXR", "YX-003": "PPP strain", "YX-004": "gpd1", "YX-005": "FDH-parent",
               "YX-006": "PPP-FDH"}
    for i, plan in enumerate(W.STRAIN_PLAN):
        s = campaign.strains[plan.strain_id]
        n = len(s["clones"])
        chosen = "" if n == 1 else (f"c{s['chosen'] + 1}" if s["chosen"] is not None else "")
        geno = plan.description if not plan.mods else "; ".join(m.label for m in plan.mods)
        if plan.fdh_nominal:
            geno = "pRS-FDH (heterologous formate dehydrogenase)"
        st = status.get(plan.strain_id, "active")
        ws.append([plan.strain_id, aliases[plan.strain_id], s["built_on"] or "-", geno, plan.date,
                   W.PEOPLE[plan.constructed_by].name, n if n > 1 else "", chosen, f"Box 3 / {chr(65 + i)}{1 + i}",
                   st, ""])
    wb.save(path)
